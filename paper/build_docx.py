"""Render the paper Markdown into an IEEE-style Word document.

Produces a single-column A4/Letter manuscript in Times New Roman with numbered
sections, formatted tables and a centred title block. Switch to the two-column
layout your target conference requires with Layout > Columns > Two in Word, or
paste the body into the official template.

Run:  python paper/build_docx.py
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

HERE = Path(__file__).parent
SOURCE = HERE / "pqc_migrator_paper.md"
OUTPUT = HERE / "PQC_Migration_Framework_Paper.docx"

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(10)


# --------------------------------------------------------------------------
# Minimal Markdown inline parsing: **bold**, *italic*, `code`
# --------------------------------------------------------------------------

INLINE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`)")


def add_runs(paragraph, text: str, *, size=BODY_SIZE, bold=False, italic=False):
    """Write ``text`` into ``paragraph``, honouring simple inline markup."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        run_bold, run_italic, mono = bold, italic, False
        if piece.startswith("**") and piece.endswith("**"):
            piece, run_bold = piece[2:-2], True
        elif piece.startswith("*") and piece.endswith("*"):
            piece, run_italic = piece[1:-1], True
        elif piece.startswith("`") and piece.endswith("`"):
            piece, mono = piece[1:-1], True
        run = paragraph.add_run(piece)
        run.bold = run_bold
        run.italic = run_italic
        run.font.size = Pt(8.5) if mono else size
        run.font.name = "Consolas" if mono else BODY_FONT


def shade(cell, hex_colour: str) -> None:
    element = OxmlElement("w:shd")
    element.set(qn("w:val"), "clear")
    element.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(element)


def style_document(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.0
    rpr = normal.element.get_or_add_rPr().get_or_add_rFonts()
    rpr.set(qn("w:eastAsia"), BODY_FONT)

    section = document.sections[0]
    for attribute in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, attribute, Pt(54))  # 0.75 inch


def title_block(document: Document, lines: list[str]) -> None:
    """Title, authors and affiliation, centred IEEE-style."""
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run(lines[0])
    run.bold = True
    run.font.size = Pt(18)
    run.font.name = BODY_FONT

    for line in lines[1:]:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(2)
        add_runs(paragraph, line, size=Pt(10.5))
    document.add_paragraph().paragraph_format.space_after = Pt(6)


def heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12 if level == 2 else 8)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = BODY_FONT
    run.font.size = Pt(11) if level == 2 else Pt(10)
    if level == 2:
        run.font.all_caps = False


def table_block(document: Document, rows: list[list[str]]) -> None:
    header, body = rows[0], rows[1:]
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, label in enumerate(header):
        cell = table.rows[0].cells[index]
        cell.text = ""
        add_runs(cell.paragraphs[0], label, size=Pt(8.5), bold=True)
        shade(cell, "E8E8E8")
    for row in body:
        cells = table.add_row().cells
        for index, value in enumerate(row[: len(header)]):
            cells[index].text = ""
            add_runs(cells[index].paragraphs[0], value, size=Pt(8.5))
    document.add_paragraph().paragraph_format.space_after = Pt(6)


def note_block(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(18)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(f"[ {text} ]")
    run.italic = True
    run.font.size = Pt(9)
    run.font.name = BODY_FONT
    run.font.color.rgb = RGBColor(0xA0, 0x30, 0x20)


def body_paragraph(document: Document, text: str, *, justify=True) -> None:
    paragraph = document.add_paragraph()
    if justify:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(5)
    add_runs(paragraph, text)


def list_item(document: Document, text: str, ordered: bool) -> None:
    paragraph = document.add_paragraph(
        style="List Number" if ordered else "List Bullet"
    )
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(3)
    add_runs(paragraph, text)


# --------------------------------------------------------------------------
# Markdown walk
# --------------------------------------------------------------------------

def build() -> Path:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    style_document(document)

    index = 0
    buffer: list[str] = []
    pending_title: list[str] = []
    in_code = False
    code_lines: list[str] = []

    def flush() -> None:
        nonlocal buffer
        if buffer:
            body_paragraph(document, " ".join(buffer).strip())
            buffer = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush()
            if in_code:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Pt(14)
                paragraph.paragraph_format.space_after = Pt(8)
                run = paragraph.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(8)
                code_lines = []
            in_code = not in_code
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if not stripped or stripped == "---":
            flush()
            index += 1
            continue

        # Tables
        if stripped.startswith("|"):
            flush()
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [c.strip() for c in lines[index].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    rows.append(cells)
                index += 1
            if rows:
                table_block(document, rows)
            continue

        # Reviewer notes
        if stripped.startswith(">"):
            flush()
            note: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                note.append(lines[index].strip().lstrip(">").strip().strip("*"))
                index += 1
            note_block(document, " ".join(n for n in note if n))
            continue

        # Headings
        if stripped.startswith("#"):
            flush()
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped.lstrip("#").strip()
            if level == 1 and not pending_title:
                pending_title.append(text)
                index += 1
                while index < len(lines) and not lines[index].strip():
                    index += 1  # skip the blank line after the title
                # the next contiguous block is the author / affiliation block
                while index < len(lines) and lines[index].strip():
                    pending_title.append(lines[index].strip())
                    index += 1
                title_block(document, pending_title)
                continue
            heading(document, text, min(level, 3))
            index += 1
            continue

        # Lists
        match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if match:
            flush()
            list_item(document, match.group(2), ordered=True)
            index += 1
            continue
        if stripped.startswith(("- ", "* ")):
            flush()
            list_item(document, stripped[2:], ordered=False)
            index += 1
            continue

        buffer.append(stripped)
        index += 1

    flush()
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(f"wrote {path} ({path.stat().st_size:,} bytes)")
