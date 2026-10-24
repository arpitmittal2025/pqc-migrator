"""Fill the Review-3 presentation template with this project's content.

Run:  python report/fill_ppt.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

import content as C

HERE = Path(__file__).parent
TEMPLATE = Path(r"C:\Users\arpit\Downloads\Review 3- 28-10- 2026.pptx")
OUTPUT = HERE / "Review_3_Team_20106UG02.pptx"
FIGURES = HERE / "figures"

REVIEW_DATE = "28/10/26"
TODO = RGBColor(0xA0, 0x30, 0x20)
INK = RGBColor(0x1A, 0x1A, 0x1A)


def shape_named(slide, *fragments, need_text=True):
    """Find a shape by name fragment; by default only one that can hold text."""
    for shape in slide.shapes:
        for fragment in fragments:
            if fragment.lower() in shape.name.lower():
                if need_text and not shape.has_text_frame:
                    continue
                return shape
    return None


def text_target(slide, *fragments, top=2.0, height=4.4, left=0.92, width=11.5):
    """A text-capable content shape, adding a textbox when the layout has none."""
    shape = shape_named(slide, *fragments)
    if shape is not None:
        return shape
    return slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height))


def place(shape, *, left, top, width, height):
    """Pin a shape's geometry in inches (all four values, see note above)."""
    shape.left = Inches(left)
    shape.top = Inches(top)
    shape.width = Inches(width)
    shape.height = Inches(height)
    return shape


def set_bullets(shape, items, *, size=16, spacing=1.0):
    """items: str, or (text, level), or (text, level, bold)."""
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for index, item in enumerate(items):
        if isinstance(item, str):
            text, level, bold = item, 0, False
        elif len(item) == 2:
            (text, level), bold = item, False
        else:
            text, level, bold = item
        para = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        para.text = text
        para.level = level
        para.space_after = Pt(6)
        para.line_spacing = spacing
        for run in para.runs:
            run.font.size = Pt(size - 2 * level)
            run.font.bold = bold
            run.font.color.rgb = INK
    return shape


def fill(presentation):
    slides = presentation.slides

    # ---------------------------------------------------------------- 1
    title = shape_named(slides[0], "Title")
    tf = title.text_frame
    tf.clear()
    first = tf.paragraphs[0]
    first.text = "B.Tech.  |  BCSE497J \u2013 Project-I"
    for run in first.runs:
        run.font.size = Pt(20)
        run.font.bold = False
    second = tf.add_paragraph()
    second.text = C.SHORT_TITLE
    for run in second.runs:
        run.font.size = Pt(30)
        run.font.bold = True

    members = shape_named(slides[0], "object 3", "TextBox")
    if members is not None:
        rows = ["Team members:"]
        rows += [f"{name}  ({reg})" for reg, name in C.TEAM]
        rows += ["", "Faculty guide:", f"{C.GUIDE_NAME}, {C.GUIDE_DESIGNATION}",
                 "", "Team ID: 20106UG02"]
        set_bullets(members, rows, size=15)

    # ---------------------------------------------------------------- 3 Aim
    set_bullets(text_target(slides[2], "Content Placeholder"), [
        ("To automate the migration of classical public-key cryptography in existing "
         "Python software to NIST-standardised post-quantum algorithms, and to prove the "
         "migrated code still works.", 0),
        ("Detect every use of RSA, ECC and Diffie-Hellman in a codebase", 1),
        ("Recommend the correct NIST replacement for each occurrence", 1),
        ("Rewrite the source automatically onto ML-KEM and ML-DSA", 1),
        ("Verify the result by executing the migrated application", 1),
        ("Report the measured cost of migration in size and time", 1),
    ], size=18)

    # ---------------------------------------------------------------- 4 Abstract
    set_bullets(text_target(slides[3], "Content Placeholder"), [
        "Shor's algorithm breaks RSA, ECC and Diffie-Hellman; NIST standardised the "
        "replacements in August 2024 (FIPS 203, 204, 205).",
        "Standardising algorithms solved only half the problem \u2014 migrating existing "
        "codebases is still slow, manual and needs cryptographic expertise.",
        "We built an end-to-end framework: AST detection \u2192 purpose inference \u2192 "
        "NIST recommendation \u2192 rule-based rewriting \u2192 functional verification \u2192 "
        "reporting.",
        "Evaluated on a 29-module labelled benchmark with 38 annotated sites: precision "
        "1.00 and recall 1.00 over 31 in-scope sites, purpose accuracy 1.00, zero false "
        "alarms on 6 clean modules.",
        "9 of 9 sites migrated across 3 sample applications; the migrated RSA and ECDSA "
        "applications execute successfully under real ML-KEM and ML-DSA.",
        "Measured cost: ML-KEM-768 keys are 4.03\u00d7 RSA-2048; ML-DSA-65 signatures are "
        "46.6\u00d7 ECDSA P-256.",
    ], size=15)

    # ---------------------------------------------------------------- 5 Lit review
    set_bullets(text_target(slides[4], "Content Placeholder", top=1.55, height=3.2), [
        ("Cryptographic discovery \u2014 NCCoE SP 1800-38: inventory is the first step, "
         "but stops at a catalogue; no recommendation, no code change.", 0),
        ("Standards and guidance \u2014 FIPS 203/204/205 and NIST white papers define the "
         "destination and crypto-agility, not a path for existing code.", 0),
        ("PQC implementations \u2014 liboqs, kyber-py, dilithium-py supply the primitives "
         "but do not migrate anything.", 0),
        ("LLM code transformation \u2014 promising on isolated functions, but output "
         "always requires human verification before it can be trusted.", 0),
    ], size=15)

    # ---------------------------------------------------------------- 6 Gap
    set_bullets(text_target(slides[5], "Content Placeholder"), [
        "Discovery tools identify cryptography but do not recommend replacements.",
        "Standards specify replacements but do not transform code.",
        "Implementations provide primitives but do not migrate anything.",
        "Transformation research changes code but does not verify the result.",
        ("No existing system connects detection \u2192 decision \u2192 transformation "
         "\u2192 verification into one workflow.", 0, True),
        ("Verification is the widest part of the gap: cryptographic correctness is not "
         "visible by inspection \u2014 it must be demonstrated by execution.", 0, True),
    ], size=16)

    # ---------------------------------------------------------------- 7 Objectives
    objectives = text_target(slides[6], "Content Placeholder")
    place(objectives, left=0.92, top=1.40, width=11.5, height=2.95)
    set_bullets(objectives, [
        "Detect RSA, ECC and DH usage with file, line, algorithm and library.",
        "Classify each occurrence as key encapsulation, key exchange or signature.",
        "Recommend a specific NIST parameter set with a written justification.",
        "Rewrite source onto PQC without modifying the originals.",
        "Verify by round-trip tests, negative tests and executing migrated code.",
        "Generate JSON, Markdown and HTML reports with a measured benchmark.",
        "Evaluate detection quantitatively against a labelled benchmark.",
    ], size=13, spacing=0.95)

    # The template's SDG and TRL reference images span the lower band, and the
    # caption ran behind them. Stack the caption above and the images below.
    sdg = shape_named(slides[6], "TextBox")
    if sdg is not None:
        place(sdg, left=0.60, top=4.45, width=12.1, height=0.85)
        set_bullets(sdg, [
            ("Identified Sustainable Development Goals:  SDG 9 \u2013 Industry, Innovation "
             "and Infrastructure (primary);  SDG 16 \u2013 Peace, Justice and Strong "
             "Institutions (secondary)", 0, True),
            ("Outcome: Conference (Scopus) \u2014 paper drafted, target venue to be "
             "finalised with the guide", 0, True),
        ], size=11.5, spacing=0.95)

    for name, left in (("Picture 7", 3.55), ("Picture 8", 7.75)):
        art = shape_named(slides[6], name, need_text=False)
        if art is not None:
            ratio = art.width / art.height
            art.height = Inches(1.45)
            art.width = int(Inches(1.45) * ratio)
            art.top = Inches(5.32)
            art.left = Inches(left)

    # ---------------------------------------------------------------- 8 Architecture
    content = text_target(slides[7], "Content Placeholder")
    content.text_frame.clear()
    # A placeholder that inherits its position from the layout loses that
    # position the moment only width/height are set, so set all four.
    place(content, left=0.92, top=1.50, width=11.5, height=0.85)
    set_bullets(content, [
        "Seven independent stages; each communicates only through explicit data records, "
        "so no stage can report an outcome an earlier stage did not produce."
    ], size=13)
    picture = FIGURES / "fig_4_1_architecture.png"
    if picture.is_file():
        width = Inches(7.0)
        slides[7].shapes.add_picture(
            str(picture), left=int((presentation.slide_width - width) / 2),
            top=Inches(2.45), width=width)

    # ---------------------------------------------------------------- 9 Functional req
    set_bullets(text_target(slides[8], "Content Placeholder"), [
        "FR1 \u2014 Recursively scan a project and report every RSA, ECC and DH usage.",
        "FR2 \u2014 Report file, line, algorithm family, library and matched construct.",
        "FR3 \u2014 Classify each occurrence as KEM, key exchange or signature.",
        "FR4 \u2014 Continue analysing files that fail to parse, via a textual fallback.",
        "FR5 \u2014 Accept target security category and signature family from config or CLI.",
        "FR6 \u2014 Recommend a NIST parameter set per finding, with a justification.",
        "FR7 \u2014 Produce migrated copies without modifying the original source.",
        "FR8 \u2014 Never report a finding as migrated unless a rule actually applied.",
        "FR9 \u2014 Verify by round-trip and negative tests, and execute migrated apps.",
        "FR10 \u2014 Generate JSON, Markdown and HTML reports with a live benchmark.",
    ], size=14)

    # ---------------------------------------------------------------- 10 Modules
    set_bullets(text_target(slides[9], "Content Placeholder"), [
        ("Detector \u2014 AST walk over every module, matched against a pattern index; "
         "alias resolution; regex fallback for unparseable files.", 0),
        ("Purpose classifier \u2014 re-infers the role of key generation from how the key "
         "is used elsewhere in the same module.", 0),
        ("Recommender \u2014 purpose selects the family, security category selects the "
         "parameter set, scoring raises it from classical strength signals.", 0),
        ("Migrator \u2014 independent rewrite rules plus a compatibility shim that keeps "
         "the classical object model over ML-KEM / ML-DSA.", 0),
        ("PQC backend \u2014 one interface over liboqs and pure-Python kyber-py / "
         "dilithium-py; unsupported algorithms raise rather than substitute.", 0),
        ("Verifier \u2014 round-trips, negative tests, and execution of migrated apps.", 0),
        ("Reporter \u2014 JSON, Markdown, HTML with a benchmark measured at run time.", 0),
    ], size=13)

    # ---------------------------------------------------------------- 11 Results
    content = text_target(slides[10], "Content Placeholder")
    place(content, left=0.92, top=1.50, width=11.5, height=2.25)
    set_bullets(content, [
        ("Detection benchmark \u2014 29 modules (18 in-scope, 5 adversarial, 6 clean), "
         "38 annotated sites.", 0, True),
        ("In-scope: precision 1.000, recall 1.000, F1 1.000 over 31 sites; purpose "
         "accuracy 1.000; 0 false alarms on 6 clean modules.", 0),
        ("Extended (adversarial): recall 0.000 of 7 \u2014 wrapper libraries, string "
         "parameters and dynamic dispatch bound syntactic analysis.", 0),
        ("Pipeline: 9/9 sites migrated across 3 apps; migrated RSA and ECDSA apps execute "
         "successfully. 70 automated tests pass.", 0),
    ], size=13)
    picture = FIGURES / "fig_7_1_sizes.png"
    if picture.is_file():
        width = Inches(6.6)
        slides[10].shapes.add_picture(
            str(picture), left=int((presentation.slide_width - width) / 2),
            top=Inches(3.90), width=width)

    # ---------------------------------------------------------------- 12 Conclusion
    set_bullets(text_target(slides[11], "Content Placeholder"), [
        ("PQC migration is a software-engineering problem with a cryptographic "
         "correctness criterion \u2014 not just an algorithm choice.", 0, True),
        ("Delivered an end-to-end verified pipeline: detection, purpose inference, "
         "recommendation, rewriting, verification and reporting.", 0),
        ("All seven objectives met; limitations measured and reported rather than "
         "concealed.", 0),
        ("Future work", 0, True),
        ("LLM-assisted transformation for the 7 adversarial sites rules cannot reach "
         "\u2014 verification stage stays, so generated code must still execute.", 1),
        ("Protocol-level restructuring of Diffie-Hellman handshakes.", 1),
        ("Wrapper-library coverage, languages beyond Python, the optimisation module, "
         "and independent evaluation on real repositories.", 1),
    ], size=14)

    # ---------------------------------------------------------------- 13 Outcome
    set_bullets(text_target(slides[12], "Content Placeholder"), [
        ("Outcome identified: International Conference Proceedings (Scopus)", 0, True),
        ("Status: full paper drafted \u2014 8 sections, 5 result tables, measured "
         "evaluation. Target venue to be finalised with the guide.", 0),
        ("Title: \u201cFrom Detection to Verified Replacement: An End-to-End Framework for "
         "Post-Quantum Cryptographic Migration of Python Codebases\u201d", 0),
        ("[Replace this slide with the submission / acceptance screenshot once the paper "
         "is submitted, and upload the proof to the given link before 27 Oct 2026.]", 0),
    ], size=14)
    for para in text_target(slides[12], "Content Placeholder").text_frame.paragraphs:
        if para.text.startswith("["):
            for run in para.runs:
                run.font.color.rgb = TODO
                run.font.italic = True

    # ---------------------------------------------------------------- 14 References
    references = text_target(slides[13], "Rectangle", "Content Placeholder")
    if references is not None:
        # The template's block starts level with the title; drop it clear.
        place(references, left=0.70, top=1.45, width=11.9, height=5.2)
        set_bullets(references, [f"[{i}] {ref.split(' [Confirm')[0]}"
                                 for i, ref in enumerate(C.REFERENCES[:8], 1)], size=11)

    # Slide 5's layout carries an empty picture placeholder; drop it so the
    # literature-review textbox is not framed by an empty content box.
    for shape in list(slides[4].shapes):
        if shape.name.startswith("Content Placeholder") and not shape.has_text_frame:
            shape._element.getparent().remove(shape._element)

    # ---------------------------------------------------------------- dates
    # The template's date placeholders are auto-updating PowerPoint fields, so
    # they show the presentation date on their own; only literal dates need setting.
    for slide in slides:
        holder = shape_named(slide, "Date Placeholder")
        if holder is None:
            continue
        for para in holder.text_frame.paragraphs:
            for run in para.runs:
                if "/" in run.text:
                    run.text = REVIEW_DATE


def build() -> Path:
    if not TEMPLATE.is_file():
        raise SystemExit(f"template not found: {TEMPLATE}")
    presentation = Presentation(str(TEMPLATE))
    fill(presentation)
    presentation.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(f"wrote {path} ({path.stat().st_size:,} bytes)")
