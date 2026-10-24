"""Fill the official BCSE497J Project-I template with this project's content.

Keeps the template's page setup, margins and section breaks intact; replaces
the placeholder text and writes the chapters, figures and tables.

Run:  python report/fill_report.py
"""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import content as C

HERE = Path(__file__).parent
TEMPLATE = Path(
    r"C:\Users\arpit\Downloads"
    r"\BCSE497J Project I Report - Template FINAlL Updated  (2) (1).docx"
)
OUTPUT = HERE / "BCSE497J_Project_I_Report_Team_20106UG02.docx"
FIGURES = HERE / "figures"

FONT = "Times New Roman"
TODO = RGBColor(0xA0, 0x30, 0x20)


# --------------------------------------------------------------------------
# low-level helpers
# --------------------------------------------------------------------------

def set_text(par, text, *, size=12, bold=False, italic=False, align=None,
             caps=False, colour=None, spacing=1.15, space_after=6):
    """Replace a paragraph's content, keeping the paragraph itself."""
    for run in list(par.runs):
        run._element.getparent().remove(run._element)
    run = par.add_run(text.upper() if caps else text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if colour is not None:
        run.font.color.rgb = colour
    rpr = run._element.get_or_add_rPr().get_or_add_rFonts()
    rpr.set(qn("w:eastAsia"), FONT)
    if align is not None:
        par.alignment = align
    par.paragraph_format.line_spacing = spacing
    par.paragraph_format.space_after = Pt(space_after)
    return par


def drop(par):
    par._element.getparent().remove(par._element)


def find(doc, needle, *, start=0):
    for i, p in enumerate(doc.paragraphs):
        if i >= start and needle.lower() in p.text.lower():
            return i, p
    return -1, None


def shade(cell, colour="E8E8E8"):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), colour)
    cell._tc.get_or_add_tcPr().append(el)


def cell_text(cell, text, *, size=10.5, bold=False, align=None):
    cell.text = ""
    par = cell.paragraphs[0]
    set_text(par, text, size=size, bold=bold, align=align, spacing=1.0, space_after=2)


# --------------------------------------------------------------------------
# body writers (append to the end of the document)
# --------------------------------------------------------------------------

class Body:
    def __init__(self, doc):
        self.doc = doc

    def para(self, text, **kw):
        return set_text(self.doc.add_paragraph(), text,
                        align=WD_ALIGN_PARAGRAPH.JUSTIFY, **kw)

    def chapter(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.page_break_before = True
        return set_text(p, text, size=14, bold=True, caps=True,
                        align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.5, space_after=6)

    def chapter_title(self, text):
        return set_text(self.doc.add_paragraph(), text, size=14, bold=True, caps=True,
                        align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.5, space_after=12)

    def h2(self, text):
        p = set_text(self.doc.add_paragraph(), text, size=13, bold=True, spacing=1.5)
        p.paragraph_format.space_before = Pt(12)
        return p

    def h3(self, text):
        p = set_text(self.doc.add_paragraph(), text, size=12, bold=True, spacing=1.5)
        p.paragraph_format.space_before = Pt(8)
        return p

    def bullets(self, items, ordered=False):
        """The template defines no list styles, so markers are set explicitly."""
        for number, item in enumerate(items, 1):
            marker = f"{number}." if ordered else "•"
            p = self.doc.add_paragraph(style="List Paragraph")
            set_text(p, f"{marker}\t{item}", align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                     space_after=5)
            p.paragraph_format.left_indent = Inches(0.65)
            p.paragraph_format.first_line_indent = Inches(-0.3)

    def table(self, caption, headers, rows, number):
        set_text(self.doc.add_paragraph(), f"Table {number}: {caption}",
                 size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                 spacing=1.0, space_after=4)
        tbl = self.doc.add_table(rows=1, cols=len(headers))
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, head in enumerate(headers):
            cell_text(tbl.rows[0].cells[i], head, size=10.5, bold=True)
            shade(tbl.rows[0].cells[i])
        for row in rows:
            cells = tbl.add_row().cells
            for i, value in enumerate(row[: len(headers)]):
                cell_text(cells[i], value, size=10.5)
        set_text(self.doc.add_paragraph(), "", size=6, space_after=6)
        return tbl

    def figure(self, filename, caption, width, number):
        path = FIGURES / filename
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        if path.is_file():
            p.add_run().add_picture(str(path), width=Inches(width))
        else:
            set_text(p, f"[missing figure: {filename}]", colour=TODO, italic=True)
        set_text(self.doc.add_paragraph(), f"Figure {number}: {caption}",
                 size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                 spacing=1.0, space_after=10)

    def code(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(text)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Consolas")


# --------------------------------------------------------------------------
# front matter
# --------------------------------------------------------------------------

def fill_front_matter(doc):
    names_line = ", ".join(name for _, name in C.TEAM)

    i, p = find(doc, "<Title of the Project>")
    set_text(p, C.TITLE, size=18, bold=True, caps=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.5)
    drop(doc.paragraphs[i + 1])  # the "(Times New Roman 18...)" hint

    i, p = find(doc, "Programme")
    if p is not None and p.text.strip() == "Programme":
        set_text(p, C.PROGRAMME, size=16, bold=True,
                 align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.5)

    i, p = find(doc, "mention names of team members")
    if p is not None:
        drop(p)

    # cover page team table
    table = doc.tables[0]
    for row, (reg, name) in zip(table.rows, C.TEAM):
        cell_text(row.cells[0], reg, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(row.cells[1], name, size=14, bold=True)

    # guide block
    guide = doc.tables[1]
    for row, value in zip(guide.rows, [C.GUIDE_NAME, C.GUIDE_DESIGNATION, C.SCHOOL]):
        bold = value == C.GUIDE_NAME
        cell_text(row.cells[0], value, size=14, bold=bold, align=WD_ALIGN_PARAGRAPH.CENTER)

    # declaration / certificate
    for needle, replacement in (
        ("<< Project Title >>", C.TITLE),
        ("<<Project Title>>", C.TITLE),
    ):
        for p in doc.paragraphs:
            if needle in p.text:
                full = p.text.replace(needle, f"\u201c{C.TITLE}\u201d")
                full = full.replace("<< Student Name (Reg. No) >>",
                                    ", ".join(f"{n} ({r})" for r, n in C.TEAM))
                full = full.replace("<<Guide Name>>", C.GUIDE_NAME)
                full = full.replace("<< Guide Name >>", C.GUIDE_NAME)
                set_text(p, full, align=WD_ALIGN_PARAGRAPH.JUSTIFY, spacing=1.5)

    for p in doc.paragraphs:
        if "<< Student Name (Reg. No) >>" in p.text:
            set_text(p, p.text.replace("<< Student Name (Reg. No) >>",
                                       ", ".join(f"{n} ({r})" for r, n in C.TEAM)),
                     align=WD_ALIGN_PARAGRAPH.JUSTIFY, spacing=1.5)

    for needle, replacement, bold in (
        ("<< Head of the Department >>", C.HOD, True),
        ("<< Programme >>", C.PROGRAMME, False),
        ("<Head of Department\u2019s Name>", C.HOD, False),
        ("<Guide\u2019s Name>", f"{C.GUIDE_NAME}, {C.GUIDE_DESIGNATION}", False),
        ("<Department Name>", C.SCHOOL, False),
    ):
        for p in doc.paragraphs:
            if needle in p.text:
                new = p.text.replace(needle, replacement)
                colour = TODO if replacement.startswith("<<") else None
                set_text(p, new, bold=bold, spacing=1.5,
                         align=p.alignment, colour=colour)

    for p in doc.paragraphs:
        if p.text.strip() == "Name of the Candidate":
            set_text(p, names_line, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
        if p.text.strip().startswith("Note: Include your project title"):
            drop(p)
        elif p.text.strip().startswith("Note: Include your project title and student"):
            drop(p)


def fill_toc(doc):
    """Fill the contents table with this report's structure."""
    toc = doc.tables[2]
    rows = [
        ("", "Abstract", "i"),
        ("", "List of Figures", "iii"),
        ("", "List of Tables", "iv"),
        ("", "List of Abbreviations", "v"),
        ("1.", "INTRODUCTION", "1"),
        ("", "1.1 Background", "1"),
        ("", "1.2 Motivations", "2"),
        ("", "1.3 Scope of the Project", "3"),
        ("2.", "PROJECT DESCRIPTION AND GOALS", "5"),
        ("", "2.1 Literature Review", "5"),
        ("", "2.2 Research Gap", "8"),
        ("", "2.3 Objectives", "9"),
        ("", "2.4 Problem Statement", "10"),
        ("", "2.5 Project Plan", "11"),
        ("3.", "TECHNICAL SPECIFICATION", "12"),
        ("", "3.1 Requirements", "12"),
        ("", "3.1.1 Functional", "12"),
        ("", "3.1.2 Non-Functional", "13"),
        ("", "3.2 Feasibility Study", "14"),
        ("", "3.2.1 Technical Feasibility", "14"),
        ("", "3.2.2 Economic Feasibility", "15"),
        ("", "3.2.3 Social Feasibility", "15"),
        ("", "3.3 System Specification", "16"),
        ("", "3.3.1 Hardware Specification", "16"),
        ("", "3.3.2 Software Specification", "16"),
        ("4.", "DESIGN APPROACH AND DETAILS", "18"),
        ("", "4.1 System Architecture", "18"),
        ("", "4.2 Design", "19"),
        ("", "4.2.1 Data Flow Diagram", "19"),
        ("", "4.2.2 Use Case Diagram", "21"),
        ("", "4.2.3 Class Diagram", "22"),
        ("", "4.2.4 Sequence Diagram", "23"),
        ("5.", "METHODOLOGY AND TESTING", "24"),
        ("", "5.1 Module Description", "24"),
        ("", "5.2 Testing", "26"),
        ("6.", "PROJECT DEMONSTRATION", "28"),
        ("7.", "RESULT AND DISCUSSION", "30"),
        ("8.", "CONCLUSION", "36"),
        ("9.", "REFERENCES", "38"),
        ("", "APPENDIX A \u2013 SAMPLE CODE", "39"),
    ]
    header = toc.rows[0]
    while len(toc.rows) > 1:
        toc._tbl.remove(toc.rows[-1]._tr)
    for sl, title, page in rows:
        cells = toc.add_row().cells
        bold = bool(sl) or title in ("Abstract",)
        cell_text(cells[0], sl, size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(cells[1], title, size=11, bold=bold)
        cell_text(cells[2], page, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    for cell in header.cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True


def fill_lists(doc):
    """List of Figures, List of Tables, Abbreviations, Symbols."""
    # figures table
    fig_tbl = doc.tables[3]
    while len(fig_tbl.rows) > 1:
        fig_tbl._tbl.remove(fig_tbl.rows[-1]._tr)
    for number, caption, page in C.FIGURES:
        cells = fig_tbl.add_row().cells
        cell_text(cells[0], number, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(cells[1], caption, size=11)
        cell_text(cells[2], str(page), size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

    # tables list: template uses tab-separated paragraphs
    i, p = find(doc, "Table No.\tTitle")
    if p is not None:
        set_text(p, "Table No.\tTitle\tPage No.", size=11, bold=True, spacing=1.5)
        anchor = p
        sample_index, sample = find(doc, "2.1\tTable caption")
        if sample is not None:
            drop(sample)
        for number, caption, page in reversed(C.TABLES):
            new = anchor._element.makeelement(qn("w:p"), {})
            anchor._element.addnext(new)
            from docx.text.paragraph import Paragraph
            set_text(Paragraph(new, anchor._parent),
                     f"{number}\t{caption}\t{page}", size=11, spacing=1.5)

    # abbreviations
    start, _ = find(doc, "List of Abbreviations")
    stop, _ = find(doc, "All the abbreviations should follow")
    if start != -1 and stop != -1:
        for p in doc.paragraphs[start + 1: stop]:
            drop(p)
        anchor = doc.paragraphs[start]
        from docx.text.paragraph import Paragraph
        for short, long in reversed(C.ABBREVIATIONS):
            new = anchor._element.makeelement(qn("w:p"), {})
            anchor._element.addnext(new)
            set_text(Paragraph(new, anchor._parent), f"{short}\t{long}",
                     size=12, spacing=1.5, space_after=2)
        _, note = find(doc, "All the abbreviations should follow")
        if note is not None:
            drop(note)

    # symbols section is not applicable to this project
    start, _ = find(doc, "Symbols and Notations")
    if start != -1:
        anchor = doc.paragraphs[start]
        for p in list(doc.paragraphs[start + 1: start + 6]):
            if p.text.strip() and "ABSTRACT" not in p.text.upper():
                drop(p)
        from docx.text.paragraph import Paragraph
        new = anchor._element.makeelement(qn("w:p"), {})
        anchor._element.addnext(new)
        set_text(Paragraph(new, anchor._parent),
                 "Not applicable. This project does not introduce new mathematical notation; "
                 "algorithm names follow the designations used in FIPS 203, FIPS 204 and "
                 "FIPS 205.", size=12, spacing=1.5,
                 align=WD_ALIGN_PARAGRAPH.JUSTIFY)


def fill_abstract(doc):
    index, heading = find(doc, "ABSTRACT")
    if heading is None:
        return
    set_text(heading, "ABSTRACT", size=16, bold=True, caps=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.5, space_after=12)
    # remove the formatting hints that follow
    for p in list(doc.paragraphs[index + 1: index + 6]):
        if p.text.strip().startswith(("(Times New Roman", "<Contents", "One page")):
            drop(p)
    from docx.text.paragraph import Paragraph
    anchor = doc.paragraphs[index]
    for block in reversed(C.ABSTRACT.strip().split("\n\n")):
        new = anchor._element.makeelement(qn("w:p"), {})
        anchor._element.addnext(new)
        set_text(Paragraph(new, anchor._parent), block.replace("\n", " "),
                 size=12, spacing=1.15, align=WD_ALIGN_PARAGRAPH.JUSTIFY)


# --------------------------------------------------------------------------
# body
# --------------------------------------------------------------------------

def clear_body(doc):
    """Delete the template's placeholder chapters, keeping section breaks sane."""
    index, _ = find(doc, "CHAPTER 1")
    if index == -1:
        return
    for p in list(doc.paragraphs[index:]):
        drop(p)


def write_body(doc):
    body = Body(doc)
    fig_no = iter([n for n, _, _ in C.FIGURES])
    tbl_no = iter([n for n, _, _ in C.TABLES])

    for block in C.BODY:
        kind = block[0]
        if kind == "h1":
            body.chapter(block[1])
        elif kind == "h1b":
            body.chapter_title(block[1])
        elif kind == "h2":
            body.h2(block[1])
        elif kind == "h3":
            body.h3(block[1])
        elif kind == "p":
            text = block[1]
            colour = TODO if text.strip().startswith("[Insert") else None
            body.para(text, colour=colour, italic=colour is not None)
        elif kind == "bul":
            body.bullets(block[1])
        elif kind == "num":
            body.bullets(block[1], ordered=True)
        elif kind == "tbl":
            body.table(block[1], block[2], block[3], next(tbl_no))
        elif kind == "fig":
            body.figure(block[1], block[2], block[3], next(fig_no))
        elif kind == "code":
            body.code(block[1])

    # ------------------------------------------------------------ references
    body.chapter("CHAPTER 9")
    body.chapter_title("REFERENCES")
    body.para("References are listed in order of recency, with the most recently published "
              "source first.", space_after=10)
    for i, ref in enumerate(C.REFERENCES, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.45)
        p.paragraph_format.first_line_indent = Inches(-0.45)
        set_text(p, f"[{i}]\t{ref}", size=12, spacing=1.15,
                 align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8)

    # ------------------------------------------------------------- appendix
    body.chapter("APPENDIX A \u2013 SAMPLE CODE")
    body.para("The following extracts show the core of each stage of the pipeline. The "
              "complete source, the three sample applications, the labelled detection "
              "benchmark and the evaluation harness are maintained in the project "
              "repository.", space_after=10)

    samples = [
        ("A.1  Detection: resolving a construct to a pattern",
         "pqc_migrator/detector/scanner.py",
         'def _canonical(self, dotted: str) -> str:\n'
         '    """Expand a local alias back to the module it was imported from."""\n'
         '    root, _, rest = dotted.partition(".")\n'
         '    origin = self.aliases.get(root)\n'
         '    if not origin:\n'
         '        return ""\n'
         '    return f"{origin}.{rest}" if rest else origin\n'
         '\n'
         'def _record(self, node: ast.AST, dotted: str) -> None:\n'
         '    pattern = P.lookup(dotted)\n'
         '    if pattern is None:\n'
         '        canonical = self._canonical(dotted)\n'
         '        pattern = P.lookup(canonical) if canonical else None\n'
         '    if pattern is None:\n'
         '        return\n'
         '    line = getattr(node, "lineno", 0)\n'
         '    self.findings.append(Finding(\n'
         '        file=self.file, line=line,\n'
         '        algorithm=pattern.algorithm, purpose=pattern.purpose,\n'
         '        matched=pattern.name,\n'
         '        library=_library_for(dotted, self.aliases, pattern.library)))'),

        ("A.2  Purpose refinement from surrounding usage",
         "pqc_migrator/detector/scanner.py",
         'def _refine_ecc_purpose(file_findings: list[Finding]) -> None:\n'
         '    """An EC key used only for agreement is key exchange, not signing."""\n'
         '    ecc = [f for f in file_findings if f.algorithm == ALG_ECC]\n'
         '    if not ecc:\n'
         '        return\n'
         '    agrees = any(f.purpose == PURPOSE_KEY_EXCHANGE for f in ecc)\n'
         '    signs = any(f.matched in {"ec.ECDSA",\n'
         '                              "Ed25519PrivateKey.generate"} for f in ecc)\n'
         '    if agrees and not signs:\n'
         '        for finding in ecc:\n'
         '            if finding.matched.startswith(("ec.generate", "ec.SECP")):\n'
         '                finding.purpose = PURPOSE_KEY_EXCHANGE'),

        ("A.3  Recommendation with scoring and justification",
         "pqc_migrator/recommender/recommend.py",
         'def recommend(finding, context=None):\n'
         '    context = context or AppContext()\n'
         '    level, reasons = _score_level(finding, context.security_level)\n'
         '\n'
         '    if finding.purpose in (PURPOSE_KEM, PURPOSE_KEY_EXCHANGE):\n'
         '        algorithm = KEM_BY_LEVEL[level]\n'
         '    elif context.prefer_signature_scheme == "SLH-DSA":\n'
         '        algorithm = SLH_DSA_BY_LEVEL[level]\n'
         '    else:\n'
         '        algorithm = ML_DSA_BY_LEVEL[level]\n'
         '\n'
         '    return Recommendation(\n'
         '        finding=finding, pqc_algorithm=algorithm,\n'
         '        nist_standard=standard_for(algorithm), nist_level=level,\n'
         '        rationale="; ".join(rationale_parts) + ".")'),

        ("A.4  The compatibility layer: RSA key transport onto ML-KEM",
         "pqc_migrator/migrator/pqc_shim.py",
         'class KEMPublicKey:\n'
         '    def encrypt(self, plaintext: bytes, padding=None) -> bytes:\n'
         '        """ML-KEM encapsulate, then AES-256-GCM the payload."""\n'
         '        ciphertext, shared_secret = backends.kem_encapsulate(\n'
         '            self.algorithm, self.key_bytes)\n'
         '        envelope = _aes_gcm_encrypt(shared_secret, plaintext)\n'
         '        return struct.pack(">I", len(ciphertext)) + ciphertext + envelope\n'
         '\n'
         'class KEMPrivateKey:\n'
         '    def decrypt(self, blob: bytes, padding=None) -> bytes:\n'
         '        (ct_len,) = struct.unpack(">I", blob[:4])\n'
         '        ciphertext = blob[4:4 + ct_len]\n'
         '        shared_secret = backends.kem_decapsulate(\n'
         '            self.algorithm, self.secret_key, ciphertext)\n'
         '        return _aes_gcm_decrypt(shared_secret, blob[4 + ct_len:])\n'
         '\n'
         '    def exchange(self, peer_public_key) -> bytes:\n'
         '        raise NotImplementedError(\n'
         '            "ML-KEM is not a drop-in for DH.exchange(): the handshake must "\n'
         '            "become an explicit encapsulate / decapsulate flow.")'),

        ("A.5  Verification, including the negative checks",
         "pqc_migrator/verifier/verify.py",
         'def verify_signature(algorithm, message=DEFAULT_MESSAGE):\n'
         '    results = {}\n'
         '    public_key, secret_key = backends.sig_keypair(algorithm)\n'
         '    signature = backends.sig_sign(algorithm, secret_key, message)\n'
         '\n'
         '    results[f"signature_roundtrip[{algorithm}]"] = \\\n'
         '        backends.sig_verify(algorithm, public_key, message, signature)\n'
         '    results[f"signature_rejects_tampered_message[{algorithm}]"] = not \\\n'
         '        backends.sig_verify(algorithm, public_key, message + b"!", signature)\n'
         '    other_public, _ = backends.sig_keypair(algorithm)\n'
         '    results[f"signature_rejects_wrong_key[{algorithm}]"] = not \\\n'
         '        backends.sig_verify(algorithm, other_public, message, signature)\n'
         '    return results'),

        ("A.6  Evaluation harness: scoring the detector",
         "evaluation/evaluate.py",
         'legitimate = detected & expected_all\n'
         'spurious   = detected - expected_all\n'
         'core_found = core & detected\n'
         'core_missed = core - detected\n'
         '\n'
         'precision       = _div(len(legitimate), len(detected))\n'
         'core_recall     = _div(len(core_found), len(core))\n'
         'extended_recall = _div(len(extended & detected), len(extended))'),
    ]

    for title, source, code in samples:
        body.h3(title)
        set_text(doc.add_paragraph(), f"Source file: {source}", size=11, italic=True,
                 spacing=1.0, space_after=4)
        body.code(code)


# --------------------------------------------------------------------------

HINTS = (
    "(times new roman",
    "<contents,",
    "one page and not exceeding",
    "in the chapters, figure caption should come",
    "note: all the abbreviations",
    "<< ieee, harvard format >>",
    "note: include your project title",
)


def strip_hints(doc):
    """Remove the template's own formatting instructions."""
    removed = 0
    for p in list(doc.paragraphs):
        text = p.text.strip().lower()
        if any(text.startswith(h) for h in HINTS):
            drop(p)
            removed += 1
    return removed


def build() -> Path:
    if not TEMPLATE.is_file():
        raise SystemExit(f"template not found: {TEMPLATE}")
    doc = Document(str(TEMPLATE))

    fill_front_matter(doc)
    fill_toc(doc)
    fill_lists(doc)
    fill_abstract(doc)
    clear_body(doc)
    write_body(doc)
    removed = strip_hints(doc)

    doc.save(OUTPUT)
    print(f"template hints removed: {removed}")
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(f"wrote {path} ({path.stat().st_size:,} bytes)")
